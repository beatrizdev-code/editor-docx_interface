from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document ()

estilo_normal = doc.styles['Normal']
estilo_normal.font.name = 'Arial'
estilo_normal.font.size = Pt (12)

secao = doc.sections[0]
secao.top_margin = Cm (2)
secao.bottom_margin = Cm(2)
secao.left_margin = Cm (2.5)
secao.right_margin = Cm (2.5)

def add_paragrafo (documento, partes, alinhamento = None, espaco_depois = 6, tamanho_fonte=12):
    p = documento.add_paragraph()
    for texto, negrito in partes:
        run = p.add_run (texto)
        run.bold = negrito
        
    if alinhamento:
        p.alignment = alinhamento
    p.paragraph_format.space_after = Pt(espaco_depois)
    run.font.size = Pt(tamanho_fonte)
    
    return p

def add_titulo (documento, texto):
    return add_paragrafo (
        documento,
        [(texto, True)],
        alinhamento= WD_ALIGN_PARAGRAPH.CENTER,
        espaco_depois=6 
    )

#Topo da página: Informações pessoais do funcionário
add_titulo(doc, 'ATESTADO PSICOLÓGICO')


add_paragrafo (doc, [
    ("Instituição atendida: ", True), ('{{INSTITUICAO}}', False), 
])

add_paragrafo (doc, [
    ('Finalidade: ', True), ('{{FINALIDADE}}', False),
])

add_paragrafo(doc, [
    ('Candidato(a) avaliado: ', True), ('{{NOME}}', False)
])

add_paragrafo (doc, [
    ('Idade: ', True),('{{IDADE}} ANOS', False),
    ('\t\tSexo: ', True), ('{{SEXO}}', False),
    ])

add_paragrafo(doc, [
    ('Estado Civil: ', True), ('{{ESTADO_CIVIL}}', False),
    ('\tData da avaliação: ', True), ('{{DATA_AVALIACAO}}', False),    
])

add_paragrafo (doc, [
    ('Escolaridade: ', True), ('{{ESCOLARIDADE}}', False)

])
doc.add_paragraph()

#Texto com as informações do candidato

add_titulo (doc, 'DESCRIÇÃO DA DEMANDA')

add_paragrafo(doc,[
    (
        "A avaliação psicológica foi realizada com o intuito de levantar os "
        "aspectos comportamentais e psicológicos da candidata para o exercício "
        "da função de {{FUNCAO}}. O processo avaliativo teve como base as "
        "seguintes competências a serem investigadas: Relacionamento "
        "interpessoal, proatividade e iniciativa, produtividade, comunicação, "
        "organização, planejamento, trabalho em equipe.", False
    )
])
doc.add_paragraph()

add_titulo(doc, 'ASPECTOS GERAIS')
add_paragrafo (doc, [('{{ASPECTOS_GERAIS}}', False)])
doc.add_paragraph()

add_titulo(doc, 'ASPECTOS DA PERSONALIDADE')

add_paragrafo (doc, [
    (
    "{{NOME_PRIMEIRO}} demonstra perfil cauteloso e criterioso na condução de "
        "suas atividades, evidenciando atenção aos detalhes e preocupação com a "
        "execução adequada das tarefas. Tende a analisar as situações antes de "
        "agir, característica que favorece maior segurança e precisão na "
        "realização de suas atribuições. Apresenta equilíbrio emocional, "
        "mantendo estabilidade comportamental diante das demandas cotidianas do "
        "ambiente profissional. Demonstra bom nível de energia, com disposição e "
        "envolvimento adequados para o cumprimento de suas responsabilidades. No "
        "âmbito interpessoal, aprecia atividades realizadas em grupo, "
        "demonstrando disponibilidade para interação, cooperação e atuação "
        "conjunta com os demais. Essa característica favorece sua integração às "
        "equipes e o desenvolvimento de relações profissionais colaborativas. "
        "Evidencia ainda respeito às normas e aos procedimentos institucionais, "
        "apresentando postura disciplinada e alinhada às orientações "
        "estabelecidas pela organização. De modo geral, {{NOME_PRIMEIRO}} "
        "apresenta perfil cauteloso, atento, colaborativo e emocionalmente "
        "equilibrado, com características favoráveis para atividades que "
        "demandem atenção, responsabilidade, trabalho em equipe e cumprimento "
        "de procedimentos. De acordo com os dados obtidos na avaliação, "
        "conclui-se que o perfil da candidata encontra-se, no momento, ",
        False,
    ),
    ('{{RESULTADO}}', True),
    (
        " para o exercício da função de {{FUNCAO}}. Deve-se considerar o "
        "caráter dinâmico e não definitivo da personalidade, com isso o perfil "
        "psicológico avaliado pode sofrer influência no contexto social.",
        False,
    ),
])

add_paragrafo(doc,[
    (
        'Este Atestado Psicológico não poderá ser utilizado para fins '
        'diferentes do apontado no item "finalidade", possui caráter sigiloso '
        'e se trata de documento extrajudicial, não deve sofrer nenhum tipo de '
        'alteração ou reprodução.', True, 
    )
], tamanho_fonte=9)

doc.add_paragraph()

#Data e asssinatura
add_paragrafo(doc,[
    ('Salvador-BA, {{DATA_ASSINATURA}}.', False)],
    alinhamento = WD_ALIGN_PARAGRAPH.CENTER,
)
doc.add_paragraph()
doc.add_paragraph()

add_paragrafo(doc,[
    ('Silvana Valverde', False)],
    alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
    espaco_depois=0,
)

add_paragrafo(
    doc,
    [('Psicóloga CRP- 03/7687', False)],
    alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
)

#salvamento do arquivo
doc.save ('modelo_template.docx')
print('Template criado com sucesso: modelo_template.docx')